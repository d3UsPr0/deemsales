from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.db import transaction
from django.utils import timezone
from datetime import datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from functools import wraps
from collections import defaultdict, OrderedDict
from django.contrib.auth.models import User
import json
from .models import PrintJob
from .models import (
    MenuCategory, MenuItem, Order, UserProfile, Expense, AuditLog
)


# ============================================================
# ROLE HELPERS
# ============================================================

def get_user_role(user):
    """Get the role of a user"""
    if user.is_superuser:
        return 'admin'
    try:
        return user.profile.role
    except:
        return 'cashier'


def role_required(*allowed_roles):
    """Decorator to restrict access by role"""
    def decorator(view_func):
        @wraps(view_func)
        @login_required
        def wrapper(request, *args, **kwargs):
            role = get_user_role(request.user)
            if role not in allowed_roles:
                messages.error(request, 'You do not have permission to access this page.')
                return redirect('home')
            return view_func(request, *args, **kwargs)
        return wrapper
    return decorator


def admin_required(view_func):
    return role_required('admin')(view_func)


def reports_required(view_func):
    return role_required('admin', 'manager')(view_func)


def expenses_required(view_func):
    return role_required('admin', 'manager', 'cashier')(view_func)


def can_view_reports(user):
    return get_user_role(user) in ['admin', 'manager']


def can_manage_menu(user):
    return get_user_role(user) in ['admin', 'manager']


def can_record_expenses(user):
    return get_user_role(user) in ['admin', 'manager', 'cashier']


def get_waiters_list():
    """Get list of all active waiter users"""
    return User.objects.filter(
        profile__role='waiter',
        is_active=True
    ).select_related('profile').order_by('first_name', 'username')


# ============================================================
# AUTH VIEWS
# ============================================================

def login_view(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        
        if user is not None:
            login(request, user)
            request.session['last_activity'] = timezone.now().timestamp()
            messages.success(request, f'Welcome, {user.username}!')
            return redirect('home')
        else:
            messages.error(request, 'Invalid credentials')
    
    if request.GET.get('timeout') == '1':
        messages.warning(request, 'Your session expired due to inactivity. Please log in again.')
    
    return render(request, 'sales/login.html')


def logout_view(request):
    logout(request)
    messages.info(request, 'You have been logged out')
    return redirect('login')


@login_required
def keep_alive(request):
    """Endpoint to keep session alive"""
    request.session['last_activity'] = timezone.now().timestamp()
    request.session.modified = True
    return JsonResponse({'status': 'ok', 'message': 'Session extended'})


# ============================================================
# HOME / ORDERS
# ============================================================

@login_required
def home_view(request):
    """Main order entry page - filters by waiter role"""
    from django.db.models import Sum
    
    user_role = get_user_role(request.user)
    is_waiter = user_role == 'waiter'
    
    categories = MenuCategory.objects.all()
    menu_items = MenuItem.objects.filter(is_available=True).select_related('category')
    
    # Waiters list — only for non-waiter roles
    if is_waiter:
        waiters = User.objects.none()
    else:
        waiters = get_waiters_list()
    
    today = timezone.now().date()
    
    # Base query — filter by waiter_user if role is waiter
    base_orders = Order.objects.all()
    if is_waiter:
        base_orders = base_orders.filter(waiter_user=request.user)
    
    today_orders = base_orders.filter(
        created_at__date=today,
        is_cancelled=False  # ← Exclude cancelled
    ).select_related('menu_item', 'waiter_user', 'created_by').order_by('-created_at')
    
    # ============================================================
    # TAB 1: LATEST ORDERS (today only)
    # ============================================================
    session_groups = defaultdict(list)
    for order in today_orders:
        session_key = order.session_id if order.session_id else f"single_{order.id}"
        session_groups[session_key].append(order)
    
    grouped_orders = []
    order_number = 1
    
    for session_key, orders in session_groups.items():
        combined_items = OrderedDict()
        for order in orders:
            item_key = order.menu_item.id
            if item_key in combined_items:
                combined_items[item_key]['quantity'] += order.quantity
                combined_items[item_key]['total_price'] += float(order.total_price)
                combined_items[item_key]['order_ids'].append(order.id)
            else:
                combined_items[item_key] = {
                    'menu_item': order.menu_item,
                    'quantity': order.quantity,
                    'unit_price': float(order.unit_price),
                    'total_price': float(order.total_price),
                    'category': order.menu_item.category,
                    'order_ids': [order.id],
                }
        
        combined_list = list(combined_items.values())
        
        all_closed = all(order.status == 'closed' for order in orders)
        any_closed = any(order.status == 'closed' for order in orders)
        if all_closed:
            payment_status = 'closed'
        elif any_closed:
            payment_status = 'partial'
        else:
            payment_status = 'open'
        
        first = orders[0]
        location_display = first.get_full_location_display()
        
        grouped_orders.append({
            'order_number': order_number,
            'reference': first.reference or f"ORD-{first.id}",
            'session_id': session_key,
            'combined_items': combined_list,
            'created_at': first.created_at,
            'waiter_user': first.waiter_user,
            'waiter_name': first.get_waiter_display_name(),
            'created_by': first.created_by,
            'location_type': first.location_type,
            'location_detail': first.location_detail,
            'location_display': location_display,
            'notes': first.notes if first.notes else '',
            'total_quantity': sum(i['quantity'] for i in combined_list),
            'total_amount': sum(i['total_price'] for i in combined_list),
            'item_count': len(combined_list),
            'is_group': len(orders) > 1 or len(combined_list) > 1,
            'first_order_id': orders[-1].group_id or orders[-1].id,
            'last_order_id': orders[0].id,
            'payment_status': payment_status,
            'is_printed': first.is_printed,
            'is_cancelled': first.is_cancelled,
            'is_locked': first.is_locked,
            'cancellation_reason': first.cancellation_reason if first.is_cancelled else '',
        })
        order_number += 1
    
    grouped_orders.sort(key=lambda x: x['created_at'], reverse=True)
    for idx, g in enumerate(grouped_orders, 1):
        g['order_number'] = idx
    
    total_collection = today_orders.aggregate(total=Sum('total_price'))['total'] or 0
    
    open_count = 0
    closed_count = 0
    for session_key, orders in session_groups.items():
        if all(order.status == 'closed' for order in orders):
            closed_count += 1
        else:
            open_count += 1
    
        # ============================================================
    # TAB 2: UNPAID ORDERS (all dates, status = open, not cancelled)
    # ============================================================
    unpaid_orders = base_orders.filter(
        status='open',
        is_cancelled=False  # ← Exclude cancelled
    ).select_related('menu_item', 'waiter_user', 'created_by').order_by('-created_at')
    
    unpaid_groups = defaultdict(list)
    for order in unpaid_orders:
        session_key = order.session_id if order.session_id else f"single_{order.id}"
        unpaid_groups[session_key].append(order)
    
    unpaid_grouped = []
    unpaid_num = 1
    
    for session_key, orders in unpaid_groups.items():
        combined_items = OrderedDict()
        for order in orders:
            item_key = order.menu_item.id
            if item_key in combined_items:
                combined_items[item_key]['quantity'] += order.quantity
                combined_items[item_key]['total_price'] += float(order.total_price)
                combined_items[item_key]['order_ids'].append(order.id)
            else:
                combined_items[item_key] = {
                    'menu_item': order.menu_item,
                    'quantity': order.quantity,
                    'unit_price': float(order.unit_price),
                    'total_price': float(order.total_price),
                    'category': order.menu_item.category,
                    'order_ids': [order.id],
                }
        
        combined_list = list(combined_items.values())
        first = orders[0]
        location_display = first.get_full_location_display()
        
        unpaid_grouped.append({
            'order_number': unpaid_num,
            'reference': first.reference or f"ORD-{first.id}",
            'session_id': session_key,
            'combined_items': combined_list,
            'created_at': first.created_at,
            'waiter_user': first.waiter_user,
            'waiter_name': first.get_waiter_display_name(),
            'created_by': first.created_by,
            'location_type': first.location_type,
            'location_detail': first.location_detail,
            'location_display': location_display,
            'notes': first.notes if first.notes else '',
            'total_quantity': sum(i['quantity'] for i in combined_list),
            'total_amount': sum(i['total_price'] for i in combined_list),
            'item_count': len(combined_list),
            'is_group': len(orders) > 1 or len(combined_list) > 1,
            'first_order_id': orders[-1].group_id or orders[-1].id,
            'last_order_id': orders[0].id,
            'days_old': (today - first.created_at.date()).days,
            'is_printed': first.is_printed,
            'is_cancelled': first.is_cancelled,
            'is_locked': first.is_locked,
            'cancellation_reason': first.cancellation_reason if first.is_cancelled else '',
        })
        unpaid_num += 1
    
    unpaid_grouped.sort(key=lambda x: x['created_at'], reverse=True)
    for idx, g in enumerate(unpaid_grouped, 1):
        g['order_number'] = idx
    
    unpaid_total = sum(g['total_amount'] for g in unpaid_grouped)
    
        # ============================================================
    # PAGINATION for both tabs (20 per page)
    # ============================================================
    from django.core.paginator import Paginator
    
    # Latest Orders pagination
    latest_page_num = request.GET.get('latest_page', 1)
    latest_paginator = Paginator(grouped_orders, 20)
    latest_page = latest_paginator.get_page(latest_page_num)
    
    # Unpaid Orders pagination
    unpaid_page_num = request.GET.get('unpaid_page', 1)
    unpaid_paginator = Paginator(unpaid_grouped, 20)
    unpaid_page = unpaid_paginator.get_page(unpaid_page_num)
    
    context = {
        'categories': categories,
        'menu_items': menu_items,
        'waiters': waiters,
        'grouped_orders': latest_page.object_list,   # Only current page
        'latest_page': latest_page,
        'latest_paginator': latest_paginator,
        'unpaid_grouped': unpaid_page.object_list,   # Only current page
        'unpaid_page': unpaid_page,
        'unpaid_paginator': unpaid_paginator,
        'unpaid_total': unpaid_total,
        'unpaid_count': len(unpaid_grouped),
        'total_collection': total_collection,
        'today': today,
        'total_orders_count': today_orders.count(),
        'total_sessions': len(grouped_orders),
        'open_count': open_count,
        'closed_count': closed_count,
        'user_role': user_role,
        'is_waiter': is_waiter,
        'can_view_reports': can_view_reports(request.user),
        'can_manage_users': get_user_role(request.user) == 'admin',
        'can_manage_menu': can_manage_menu(request.user),
        'can_record_expenses': can_record_expenses(request.user),
    }
    
    return render(request, 'sales/home.html', context)


@login_required
def submit_orders(request):
    """Handle order submission - auto-assigns waiter for waiter role"""
    if request.method == 'POST':
        try:
            with transaction.atomic():
                session_id = datetime.now().strftime('%Y%m%d%H%M%S') + str(request.user.id)
                
                # Generate readable reference: ORD-XXXX
                from .models import Order as OrderModel
                last_ref = OrderModel.objects.filter(reference__startswith='ORD-').order_by('-id').first()
                if last_ref and last_ref.reference:
                    try:
                        last_num = int(last_ref.reference.replace('ORD-', ''))
                    except:
                        last_num = 0
                else:
                    last_num = 0
                new_ref_num = last_num + 1
                reference = f"ORD-{new_ref_num}"
                
                user_role = get_user_role(request.user)
                
                # Waiter assignment logic
                if user_role == 'waiter':
                    # Waiter → auto-assign to self
                    waiter_user = request.user
                else:
                    # Admin/Manager/Cashier → use selected waiter or self
                    waiter_id = request.POST.get('waiter')
                    if waiter_id == 'self' or not waiter_id:
                        waiter_user = request.user
                    else:
                        try:
                            waiter_user = User.objects.get(id=waiter_id)
                        except User.DoesNotExist:
                            waiter_user = request.user
                
                notes = request.POST.get('notes', '')
                location_type = request.POST.get('location_type', 'table')
                location_detail = request.POST.get('location_detail', '')
                
                menu_items = MenuItem.objects.filter(is_available=True)
                orders_created = 0
                first_order_id = None
                
                for item in menu_items:
                    is_checked = request.POST.get(f'item_{item.id}')
                    quantity_str = request.POST.get(f'quantity_{item.id}', '0')
                    
                    if is_checked == 'on':
                        quantity = int(quantity_str) if quantity_str else 1
                        
                        if quantity > 0:
                            for _ in range(quantity):
                                order = Order.objects.create(
                                    menu_item=item,
                                    quantity=1,
                                    unit_price=item.price,
                                    total_price=item.price,
                                    waiter_user=waiter_user,
                                    location_type=location_type,
                                    location_detail=location_detail if location_detail else None,
                                    notes=notes,
                                    created_by=request.user,
                                    session_id=session_id
                                )
                                # Set group_id — first order of the group becomes the stable reference
                                if first_order_id is None:
                                    first_order_id = order.id
                                order.group_id = first_order_id
                                order.save(update_fields=['group_id'])
                                orders_created += 1
                
                if orders_created > 0:
                    messages.success(request, f'{orders_created} orders recorded successfully!')
                else:
                    messages.warning(request, 'No items selected. Please select at least one item.')
                
                return redirect('home')
                
        except Exception as e:
            messages.error(request, f'Error saving orders: {str(e)}')
            return redirect('home')
    
    return redirect('home')


# ============================================================
# ORDER EDIT / DELETE
# ============================================================

@login_required
def edit_order(request, order_id):
    """Get order data for editing"""
    if request.method == 'GET':
        order = get_object_or_404(Order, id=order_id)
        data = {
            'id': order.id,
            'menu_item_id': order.menu_item.id,
            'menu_item_name': order.menu_item.name,
            'quantity': order.quantity,
            'waiter_id': order.waiter_user.id if order.waiter_user else None,
            'waiter_name': order.get_waiter_display_name(),
            'location_type': order.location_type,
            'location_detail': order.location_detail or '',
            'notes': order.notes or '',
        }
        return JsonResponse(data)


@login_required
def update_order(request, order_id):
    """Update an existing order (with audit logging)"""
    if request.method == 'POST':
        try:
            order = get_object_or_404(Order, id=order_id)
            
            from .audit import serialize_instance, log_action
            data_before = serialize_instance(order)
            
            menu_item_id = request.POST.get('menu_item')
            quantity = int(request.POST.get('quantity', 1))
            waiter_id = request.POST.get('waiter')
            location_type = request.POST.get('location_type', 'table')
            location_detail = request.POST.get('location_detail', '')
            notes = request.POST.get('notes', '')
            
            order.menu_item = MenuItem.objects.get(id=menu_item_id)
            order.quantity = quantity
            order.total_price = order.menu_item.price * quantity
            order.unit_price = order.menu_item.price
            
            # Waiter assignment
            if waiter_id == 'self' or not waiter_id:
                order.waiter_user = request.user
            else:
                try:
                    order.waiter_user = User.objects.get(id=waiter_id)
                except User.DoesNotExist:
                    order.waiter_user = request.user
            
            order.location_type = location_type
            order.location_detail = location_detail if location_detail else None
            order.notes = notes
            order.save()
            
            log_action(request, 'update', order, data_before=data_before)
            
            messages.success(request, 'Order updated successfully!')
            return redirect('home')
            
        except Exception as e:
            messages.error(request, f'Error updating order: {str(e)}')
            return redirect('home')
    
    return redirect('home')

@login_required
def quick_edit_qty(request):
    """Change quantity of one item — or remove it if qty=0"""
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    
    try:
        session_id = request.POST.get('session_id')
        item_id = int(request.POST.get('item_id'))
        new_qty = int(request.POST.get('new_qty', 1))
        
        orders = Order.objects.filter(session_id=session_id, menu_item_id=item_id)
        
        if not orders.exists():
            return JsonResponse({'error': 'Item not found'}, status=404)
        
        # LOCK CHECK
        first = orders.first()
        if first.is_printed:
            return JsonResponse({'error': 'Cannot modify a printed order.'}, status=400)
        if first.is_cancelled:
            return JsonResponse({'error': 'Cannot modify a cancelled order.'}, status=400)
        
        # PAID CHECK
        if all(o.status == 'closed' for o in orders):
            return JsonResponse({'error': 'Cannot modify a PAID order. Mark it as UNPAID first.'}, status=400)
        
        if new_qty <= 0:
            orders.delete()
            remaining = Order.objects.filter(session_id=session_id)
            if not remaining.exists():
                return JsonResponse({'success': True, 'message': 'Item removed. Order is now empty.'})
            return JsonResponse({'success': True, 'message': 'Item removed'})
        
        first.quantity = new_qty
        first.total_price = first.unit_price * new_qty
        first.save()
        
        for extra in orders[1:]:
            extra.delete()
        
        return JsonResponse({'success': True, 'message': f'Quantity updated to {new_qty}'})
    
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@login_required
def delete_order(request, order_id):
    """Delete an order (with audit logging)"""
    if request.method == 'POST':
        order = get_object_or_404(Order, id=order_id)
        
        # LOCK CHECK
        if order.is_printed:
            messages.error(request, 'Cannot delete a printed order. Use Cancel instead.')
            return redirect('home')
        
        if order.is_cancelled:
            messages.error(request, 'Cannot delete a cancelled order.')
            return redirect('home')
        
        from .audit import serialize_instance, log_action
        data_before = serialize_instance(order)
        log_action(request, 'delete', order, data_before=data_before)
        
        order.delete()
        messages.success(request, 'Order deleted successfully!')
    return redirect('home')


@login_required
def delete_group_orders(request, session_id):
    """Delete all orders in a group session"""
    if request.method == 'POST':
        orders = Order.objects.filter(session_id=session_id)
        count = orders.count()
        
        if count == 0:
            messages.error(request, 'Order not found.')
            return redirect('home')
        
        # LOCK CHECK
        if orders.filter(is_printed=True).exists():
            messages.error(request, 'Cannot delete a printed order. Use Cancel instead.')
            return redirect('home')
        
        if orders.filter(is_cancelled=True).exists():
            messages.error(request, 'Cannot delete a cancelled order.')
            return redirect('home')
        
        from .audit import serialize_instance, log_action
        for order in orders:
            data_before = serialize_instance(order)
            log_action(request, 'delete', order, data_before=data_before, 
                       notes=f'Group delete ({count} items)')
        
        orders.delete()
        messages.success(request, f'Deleted {count} orders from the group.')
    return redirect('home')


@login_required
def view_group_orders(request, session_id):
    """Get details of a group order"""
    if request.method == 'GET':
        orders = Order.objects.filter(session_id=session_id).select_related('menu_item', 'waiter_user')
        
        orders_data = []
        total = 0
        for order in orders:
            orders_data.append({
                'id': order.id,
                'menu_item': order.menu_item.name,
                'category': order.menu_item.category.name,
                'quantity': order.quantity,
                'unit_price': float(order.unit_price),
                'total_price': float(order.total_price),
            })
            total += float(order.total_price)
        
        return JsonResponse({
            'orders': orders_data,
            'total': total,
            'waiter': orders[0].get_waiter_display_name() if orders else 'N/A',
            'notes': orders[0].notes if orders else '',
        })


@login_required
def get_order_group_info(request, session_id):
    """Get information about an order group for adding items"""
    if request.method == 'GET':
        orders = Order.objects.filter(session_id=session_id)
        
        if not orders.exists():
            return JsonResponse({'error': 'Order not found'}, status=404)
        
        first_order = orders.first()
        
        combined_items = OrderedDict()
        for order in orders:
            item_key = order.menu_item.id
            if item_key in combined_items:
                combined_items[item_key]['quantity'] += order.quantity
            else:
                combined_items[item_key] = {
                    'menu_item_id': order.menu_item.id,
                    'menu_item_name': order.menu_item.name,
                    'quantity': order.quantity,
                    'unit_price': float(order.unit_price),
                }
        
        return JsonResponse({
            'session_id': session_id,
            'waiter_id': first_order.waiter_user.id if first_order.waiter_user else None,
            'waiter_name': first_order.get_waiter_display_name(),
            'location_type': first_order.location_type,
            'location_detail': first_order.location_detail or '',
            'notes': first_order.notes or '',
            'created_at': first_order.created_at.strftime('%H:%M'),
            'total_items': orders.count(),
            'total_amount': sum(float(o.total_price) for o in orders),
            'existing_items': list(combined_items.values()),
        })


@login_required
def add_to_order(request, session_id):
    """Add new items to an existing order — merging with existing quantities"""
    if request.method == 'POST':
        try:
            with transaction.atomic():
                existing_orders = Order.objects.filter(session_id=session_id)
                existing_order = existing_orders.first()
                
                if not existing_order:
                    messages.error(request, 'Order not found!')
                    return redirect('home')
                
                # LOCK CHECK
                if existing_order.is_printed:
                    messages.error(request, 'Cannot add items to a printed order.')
                    return redirect('home')
                
                if existing_order.is_cancelled:
                    messages.error(request, 'Cannot add items to a cancelled order.')
                    return redirect('home')
                
                # PAID CHECK — cannot add to paid orders
                if all(o.status == 'closed' for o in existing_orders):
                    messages.error(request, 'Cannot add items to a PAID order. Mark it as UNPAID first.')
                    return redirect('home')
                
                existing_by_item = defaultdict(list)
                for o in existing_orders:
                    existing_by_item[o.menu_item_id].append(o)
                
                menu_items = MenuItem.objects.filter(is_available=True)
                items_added = 0
                
                for item in menu_items:
                    is_checked = request.POST.get(f'item_{item.id}')
                    quantity_str = request.POST.get(f'quantity_{item.id}', '0')
                    
                    if is_checked == 'on':
                        quantity = int(quantity_str) if quantity_str else 1
                        
                        if quantity > 0:
                            if item.id in existing_by_item:
                                first = existing_by_item[item.id][0]
                                new_qty = first.quantity + quantity
                                first.quantity = new_qty
                                first.total_price = first.unit_price * new_qty
                                first.save()
                                
                                for extra in existing_by_item[item.id][1:]:
                                    extra.delete()
                            else:
                                # NEW item: create individual rows
                                group_id = existing_order.group_id or existing_order.id
                                for _ in range(quantity):
                                    Order.objects.create(
                                        menu_item=item,
                                        quantity=1,
                                        unit_price=item.price,
                                        total_price=item.price,
                                        waiter_user=existing_order.waiter_user,
                                        location_type=existing_order.location_type,
                                        location_detail=existing_order.location_detail,
                                        notes=existing_order.notes,
                                        status=existing_order.status,
                                        created_by=request.user,
                                        session_id=session_id,
                                        group_id=group_id
                                    )
                            items_added += quantity
                
                if items_added > 0:
                    messages.success(request, f'{items_added} item(s) added to order successfully!')
                else:
                    messages.warning(request, 'No items selected to add.')
                
                return redirect('home')
                
        except Exception as e:
            messages.error(request, f'Error adding items: {str(e)}')
            return redirect('home')
    
    return redirect('home')


@login_required
def toggle_group_status(request, session_id):
    """Toggle Open/Closed status for entire order group"""
    if request.method == 'POST':
        orders = Order.objects.filter(session_id=session_id)
        if orders.exists():
            all_closed = all(order.status == 'closed' for order in orders)
            
            if all_closed:
                orders.update(status='open', closed_at=None)
                status_msg = f'Order marked as UNPAID ({orders.count()} items)'
            else:
                orders.update(status='closed', closed_at=timezone.now())
                status_msg = f'Order marked as PAID ({orders.count()} items)'
            
            if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
                return JsonResponse({'success': True, 'message': status_msg})
            
            messages.success(request, status_msg)
    
    return redirect('home')


@login_required
def get_order_group_full(request, session_id):
    """Get full order data for editing"""
    if request.method == 'GET':
        orders = Order.objects.filter(session_id=session_id).select_related('menu_item', 'waiter_user')
        
        if not orders.exists():
            return JsonResponse({'error': 'Order not found'}, status=404)
        
        first = orders.first()
        
        combined = OrderedDict()
        for o in orders:
            key = o.menu_item.id
            if key in combined:
                combined[key]['quantity'] += o.quantity
                combined[key]['total_price'] += float(o.total_price)
                combined[key]['order_ids'].append(o.id)
            else:
                combined[key] = {
                    'menu_item_id': o.menu_item.id,
                    'menu_item_name': o.menu_item.name,
                    'category': o.menu_item.category.name,
                    'unit_price': float(o.unit_price),
                    'quantity': o.quantity,
                    'total_price': float(o.total_price),
                    'order_ids': [o.id],
                }
        
        items = list(combined.values())
        
        all_closed = all(o.status == 'closed' for o in orders)
        any_closed = any(o.status == 'closed' for o in orders)
        status = 'closed' if all_closed else ('mixed' if any_closed else 'open')
        
        return JsonResponse({
            'session_id': session_id,
            'items': items,
            'waiter_id': first.waiter_user.id if first.waiter_user else None,
            'waiter_name': first.get_waiter_display_name(),
            'location_type': first.location_type,
            'location_detail': first.location_detail or '',
            'notes': first.notes or '',
            'status': status,
            'total_amount': sum(i['total_price'] for i in items),
            'total_quantity': sum(i['quantity'] for i in items),
        })
    
    return JsonResponse({'error': 'Invalid method'}, status=405)


@login_required
def update_order_group(request, session_id):
    """Update full order — replace items, update quantities, change waiter/location/notes"""
    if request.method == 'POST':
        try:
            with transaction.atomic():
                existing_orders = Order.objects.filter(session_id=session_id)
                
                if not existing_orders.exists():
                    messages.error(request, 'Order not found!')
                    return redirect('home')
                
                # LOCK CHECK
                if existing_orders.filter(is_printed=True).exists():
                    messages.error(request, 'Cannot modify a printed order. Use cancel/reprint.')
                    return redirect('home')
                
                if existing_orders.filter(is_cancelled=True).exists():
                    messages.error(request, 'Cannot modify a cancelled order.')
                    return redirect('home')
                
                # PAID CHECK — cannot modify paid orders
                if all(o.status == 'closed' for o in existing_orders):
                    messages.error(request, 'Cannot modify a PAID order. Mark it as UNPAID first.')
                    return redirect('home')
                
                first = existing_orders.first()
                
                # Waiter assignment
                waiter_id = request.POST.get('waiter')
                if waiter_id == 'self' or not waiter_id:
                    new_waiter_user = request.user
                else:
                    try:
                        new_waiter_user = User.objects.get(id=waiter_id)
                    except User.DoesNotExist:
                        new_waiter_user = request.user
                
                location_type = request.POST.get('location_type', 'table')
                location_detail = request.POST.get('location_detail', '')
                notes = request.POST.get('notes', '')
                
                # Parse items
                updated_items = {}
                for key, value in request.POST.items():
                    if key.startswith('item_qty_'):
                        menu_item_id = int(key.replace('item_qty_', ''))
                        try:
                            qty = int(value)
                        except:
                            qty = 0
                        if qty > 0:
                            updated_items[menu_item_id] = qty
                
                if not updated_items:
                    messages.error(request, 'Order must have at least one item.')
                    return redirect('home')
                
                existing_item_ids = set(existing_orders.values_list('menu_item_id', flat=True))
                
                # DELETE items no longer in order
                for item_id in existing_item_ids:
                    if item_id not in updated_items:
                        existing_orders.filter(menu_item_id=item_id).delete()
                
                # UPDATE or CREATE items
                existing_by_item = defaultdict(list)
                for o in existing_orders:
                    existing_by_item[o.menu_item_id].append(o)
                
                for menu_item_id, new_qty in updated_items.items():
                    try:
                        menu_item = MenuItem.objects.get(id=menu_item_id)
                    except MenuItem.DoesNotExist:
                        continue
                    
                    if menu_item_id in existing_by_item:
                        rows = existing_by_item[menu_item_id]
                        first_row = rows[0]
                        first_row.quantity = new_qty
                        first_row.unit_price = menu_item.price
                        first_row.total_price = menu_item.price * new_qty
                        first_row.waiter_user = new_waiter_user
                        first_row.location_type = location_type
                        first_row.location_detail = location_detail or None
                        first_row.notes = notes
                        first_row.save()
                        
                        for extra in rows[1:]:
                            extra.delete()
                    else:
                        Order.objects.create(
                            menu_item=menu_item,
                            quantity=new_qty,
                            unit_price=menu_item.price,
                            total_price=menu_item.price * new_qty,
                            waiter_user=new_waiter_user,
                            location_type=location_type,
                            location_detail=location_detail or None,
                            notes=notes,
                            status=first.status,
                            created_by=request.user,
                            session_id=session_id,
                            group_id=first.group_id
                        )
                
                # Update on all remaining rows
                Order.objects.filter(session_id=session_id).update(
                    waiter_user=new_waiter_user,
                    location_type=location_type,
                    location_detail=location_detail or None,
                    notes=notes,
                )
                
                messages.success(request, 'Order updated successfully!')
                return redirect('home')
                
        except Exception as e:
            messages.error(request, f'Error updating order: {str(e)}')
            return redirect('home')
    
    return redirect('home')


# ============================================================
# EXCEL EXPORT
# ============================================================

@reports_required
def export_excel(request):
    """Export orders, expenses, and item sales to Excel"""
    today = timezone.now().date()
    date_from = None
    date_to = None
    label = ""
    
    preset = request.GET.get('preset')
    
    if preset == 'today':
        date_from = today
        date_to = today
        label = f"Today ({today})"
    elif preset == 'yesterday':
        date_from = today - timedelta(days=1)
        date_to = today - timedelta(days=1)
        label = f"Yesterday ({date_from})"
    elif preset == 'this_week':
        date_from = today - timedelta(days=today.weekday())
        date_to = today
        label = f"This Week ({date_from} to {date_to})"
    elif preset == 'this_month':
        date_from = today.replace(day=1)
        date_to = today
        label = f"This Month ({date_from} to {date_to})"
    elif preset == 'last_month':
        first_of_this_month = today.replace(day=1)
        last_day_of_last_month = first_of_this_month - timedelta(days=1)
        date_from = last_day_of_last_month.replace(day=1)
        date_to = last_day_of_last_month
        label = f"Last Month ({date_from} to {date_to})"
    else:
        date_from_str = request.GET.get('date_from')
        date_to_str = request.GET.get('date_to')
        
        if date_from_str:
            date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date()
        if date_to_str:
            date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date()
        
        if date_from and date_to:
            label = f"{date_from} to {date_to}"
        elif date_from:
            label = f"From {date_from}"
            date_to = today
        elif date_to:
            label = f"Until {date_to}"
            date_from = today - timedelta(days=30)
        else:
            date_from = today
            date_to = today
            label = f"Today ({today})"
    
    # Exclude cancelled — they go to their own sheet
    orders = Order.objects.filter(
        created_at__date__gte=date_from,
        created_at__date__lte=date_to,
        is_cancelled=False
    ).select_related('menu_item', 'waiter_user', 'created_by').order_by('-created_at')
    
    expenses = Expense.objects.filter(
        date__gte=date_from,
        date__lte=date_to
    ).select_related('recorded_by').order_by('date', '-created_at')
    
    wb = Workbook()
    
    # Styling
    primary_fill = PatternFill(start_color="540b02", end_color="540b02", fill_type="solid")
    secondary_fill = PatternFill(start_color="00b652", end_color="00b652", fill_type="solid")
    accent_fill = PatternFill(start_color="f7941d", end_color="f7941d", fill_type="solid")
    danger_fill = PatternFill(start_color="dc3545", end_color="dc3545", fill_type="solid")
    success_fill = PatternFill(start_color="28a745", end_color="28a745", fill_type="solid")
    white_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    # ===== SHEET 1: ORDERS =====
    ws = wb.active
    ws.title = "Orders"
    
    ws.merge_cells('A1:K1')
    ws['A1'].value = f"DEEM Restaurant - Orders Report ({label})"
    ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws['A1'].fill = primary_fill
    ws['A1'].alignment = Alignment(horizontal='center')
    
    headers = ['Order #', 'Menu Items', 'Total Items', 'Total Qty', 'Total Amount (TZS)', 
               'Waiter', 'Location', 'Status', 'Notes', 'Created By', 'Time']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=header)
        cell.font = white_font
        cell.fill = secondary_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')
    
    session_groups = defaultdict(list)
    for order in orders:
        session_key = order.session_id if order.session_id else f"single_{order.id}"
        session_groups[session_key].append(order)
    
    row = 4
    grand_total = 0
    
    for session_key, group_orders in session_groups.items():
        group_orders.sort(key=lambda x: x.created_at)
        
        combined = OrderedDict()
        for o in group_orders:
            key = o.menu_item.name
            combined[key] = combined.get(key, 0) + o.quantity
        
        menu_items_str = ", ".join([f"{name} x{qty}" if qty > 1 else name for name, qty in combined.items()])
        total_qty = sum(o.quantity for o in group_orders)
        total_amount = sum(float(o.total_price) for o in group_orders)
        grand_total += total_amount
        
        first_order = group_orders[0]
        
        all_closed = all(o.status == 'closed' for o in group_orders)
        any_closed = any(o.status == 'closed' for o in group_orders)
        status = 'Paid' if all_closed else ('Mixed' if any_closed else 'Unpaid')
        
        location = first_order.get_full_location_display()
        waiter_name = first_order.get_waiter_display_name()
        
        first_order = group_orders[0]
        ws.cell(row=row, column=1, value=f"#{first_order.group_id or first_order.id}").border = border
        ws.cell(row=row, column=2, value=menu_items_str).border = border
        ws.cell(row=row, column=3, value=len(combined)).border = border
        ws.cell(row=row, column=4, value=total_qty).border = border
        ws.cell(row=row, column=5, value=total_amount).border = border
        ws.cell(row=row, column=6, value=waiter_name).border = border
        ws.cell(row=row, column=7, value=location).border = border
        ws.cell(row=row, column=8, value=status).border = border
        ws.cell(row=row, column=9, value=first_order.notes or '').border = border
        ws.cell(row=row, column=10, value=first_order.created_by.username).border = border
        ws.cell(row=row, column=11, value=first_order.created_at.strftime('%H:%M')).border = border
        
        status_cell = ws.cell(row=row, column=8)
        if status == 'Paid':
            status_cell.fill = PatternFill(start_color="c6efce", end_color="c6efce", fill_type="solid")
            status_cell.font = Font(color="006100")
        elif status == 'Unpaid':
            status_cell.fill = PatternFill(start_color="ffc7ce", end_color="ffc7ce", fill_type="solid")
            status_cell.font = Font(color="9c0006")
        else:
            status_cell.fill = PatternFill(start_color="ffeb9c", end_color="ffeb9c", fill_type="solid")
        status_cell.font = Font(color="9c6500")
        
        row += 1
    
    # Orders summary
    summary_row = row
    ws.merge_cells(f'A{summary_row}:D{summary_row}')
    summary_cell = ws.cell(row=summary_row, column=1, value="TOTAL SALES")
    summary_cell.font = Font(bold=True, size=12)
    summary_cell.fill = accent_fill
    summary_cell.border = border
    
    total_cell = ws.cell(row=summary_row, column=5, value=grand_total)
    total_cell.font = Font(bold=True, size=12, color="540b02")
    total_cell.fill = accent_fill
    total_cell.border = border
    
    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 20
    ws.column_dimensions['B'].width = 45
    
    # ===== SHEET 2: EXPENSES =====
    ws2 = wb.create_sheet("Expenses")
    
    ws2.merge_cells('A1:F1')
    ws2['A1'].value = f"DEEM Restaurant - Expenses Report ({label})"
    ws2['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws2['A1'].fill = primary_fill
    ws2['A1'].alignment = Alignment(horizontal='center')
    
    headers2 = ['Date', 'Category', 'Description', 'Receipt #', 'Amount (TZS)', 'Recorded By']
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col, value=header)
        cell.font = white_font
        cell.fill = danger_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')
    
    expense_total = 0
    expense_row = 4
    for expense in expenses:
        ws2.cell(row=expense_row, column=1, value=expense.date.strftime('%d/%m/%Y')).border = border
        ws2.cell(row=expense_row, column=2, value=expense.get_category_display()).border = border
        ws2.cell(row=expense_row, column=3, value=expense.description).border = border
        ws2.cell(row=expense_row, column=4, value=expense.receipt_ref or '-').border = border
        ws2.cell(row=expense_row, column=5, value=float(expense.amount)).border = border
        ws2.cell(row=expense_row, column=6, value=expense.recorded_by.username).border = border
        expense_total += float(expense.amount)
        expense_row += 1
    
    if expenses:
        exp_summary_row = expense_row
        ws2.merge_cells(f'A{exp_summary_row}:D{exp_summary_row}')
        ws2.cell(row=exp_summary_row, column=1, value="TOTAL EXPENSES").font = Font(bold=True, size=12, color="FFFFFF")
        ws2.cell(row=exp_summary_row, column=1).fill = danger_fill
        ws2.cell(row=exp_summary_row, column=1).border = border
        ws2.cell(row=exp_summary_row, column=5, value=expense_total).font = Font(bold=True, size=12, color="FFFFFF")
        ws2.cell(row=exp_summary_row, column=5).fill = danger_fill
        ws2.cell(row=exp_summary_row, column=5).border = border
    
    for col in range(1, 7):
        ws2.column_dimensions[get_column_letter(col)].width = 22
    ws2.column_dimensions['C'].width = 40
    
    # ===== SHEET 3: ITEM SALES + PROFIT/LOSS =====
    ws3 = wb.create_sheet("Item Sales")
    
    ws3.merge_cells('A1:F1')
    ws3['A1'].value = f"DEEM Restaurant - Item Sales Summary ({label})"
    ws3['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws3['A1'].fill = primary_fill
    ws3['A1'].alignment = Alignment(horizontal='center')
    
    headers3 = ['Menu Item', 'Category', 'Unit Price (TZS)', 'Qty Sold', 'Total Sales (TZS)', '% of Sales']
    for col, header in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=col, value=header)
        cell.font = white_font
        cell.fill = secondary_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')
    
    item_sales = {}
    for order in orders:
        key = order.menu_item.id
        if key not in item_sales:
            item_sales[key] = {
                'name': order.menu_item.name,
                'category': order.menu_item.category.name,
                'unit_price': float(order.menu_item.price),
                'qty': 0,
                'total': 0,
            }
        item_sales[key]['qty'] += order.quantity
        item_sales[key]['total'] += float(order.total_price)
    
    sorted_items = sorted(item_sales.values(), key=lambda x: x['total'], reverse=True)
    
    item_row = 4
    for item in sorted_items:
        percent = (item['total'] / grand_total * 100) if grand_total > 0 else 0
        ws3.cell(row=item_row, column=1, value=item['name']).border = border
        ws3.cell(row=item_row, column=2, value=item['category']).border = border
        ws3.cell(row=item_row, column=3, value=item['unit_price']).border = border
        ws3.cell(row=item_row, column=4, value=item['qty']).border = border
        ws3.cell(row=item_row, column=5, value=item['total']).border = border
        ws3.cell(row=item_row, column=6, value=f"{percent:.1f}%").border = border
        item_row += 1
    
    total_row = item_row
    ws3.merge_cells(f'A{total_row}:C{total_row}')
    ws3.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)
    ws3.cell(row=total_row, column=1).fill = accent_fill
    ws3.cell(row=total_row, column=1).border = border
    ws3.cell(row=total_row, column=4, value=sum(i['qty'] for i in sorted_items)).font = Font(bold=True)
    ws3.cell(row=total_row, column=4).fill = accent_fill
    ws3.cell(row=total_row, column=4).border = border
    ws3.cell(row=total_row, column=5, value=grand_total).font = Font(bold=True)
    ws3.cell(row=total_row, column=5).fill = accent_fill
    ws3.cell(row=total_row, column=5).border = border
    ws3.cell(row=total_row, column=6, value="100%").font = Font(bold=True)
    ws3.cell(row=total_row, column=6).fill = accent_fill
    ws3.cell(row=total_row, column=6).border = border
    
    # Profit/Loss
    profit_start = total_row + 3
    ws3.merge_cells(f'A{profit_start}:F{profit_start}')
    pl_title = ws3.cell(row=profit_start, column=1, value="PROFIT / LOSS SUMMARY")
    pl_title.font = Font(size=13, bold=True, color="FFFFFF")
    pl_title.fill = primary_fill
    pl_title.alignment = Alignment(horizontal='center')
    pl_title.border = border
    
    ws3.merge_cells(f'A{profit_start+1}:D{profit_start+1}')
    ws3.cell(row=profit_start+1, column=1, value="Total Sales").font = Font(bold=True)
    ws3.cell(row=profit_start+1, column=1).border = border
    ws3.cell(row=profit_start+1, column=1).fill = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")
    sales_cell = ws3.cell(row=profit_start+1, column=5, value=grand_total)
    sales_cell.font = Font(bold=True, color="006100")
    sales_cell.fill = PatternFill(start_color="e8f5e9", end_color="e8f5e9", fill_type="solid")
    sales_cell.border = border
    
    ws3.merge_cells(f'A{profit_start+2}:D{profit_start+2}')
    ws3.cell(row=profit_start+2, column=1, value="Total Expenses").font = Font(bold=True)
    ws3.cell(row=profit_start+2, column=1).border = border
    ws3.cell(row=profit_start+2, column=1).fill = PatternFill(start_color="ffebee", end_color="ffebee", fill_type="solid")
    exp_cell = ws3.cell(row=profit_start+2, column=5, value=expense_total)
    exp_cell.font = Font(bold=True, color="9c0006")
    exp_cell.fill = PatternFill(start_color="ffebee", end_color="ffebee", fill_type="solid")
    exp_cell.border = border
    
    net = grand_total - expense_total
    ws3.merge_cells(f'A{profit_start+3}:D{profit_start+3}')
    net_label = ws3.cell(row=profit_start+3, column=1, value="NET PROFIT / LOSS")
    net_label.font = Font(bold=True, size=12, color="FFFFFF")
    net_label.fill = success_fill if net >= 0 else danger_fill
    net_label.border = border
    
    net_cell = ws3.cell(row=profit_start+3, column=5, value=net)
    net_cell.font = Font(bold=True, size=12, color="FFFFFF")
    net_cell.fill = success_fill if net >= 0 else danger_fill
    net_cell.border = border
    
    ws3.merge_cells(f'A{profit_start+4}:D{profit_start+4}')
    cash_label = ws3.cell(row=profit_start+4, column=1, value="CASH ON HAND")
    cash_label.font = Font(bold=True, color="FFFFFF")
    cash_label.fill = accent_fill
    cash_label.border = border
    cash_cell = ws3.cell(row=profit_start+4, column=5, value=net)
    cash_cell.font = Font(bold=True, color="FFFFFF")
    cash_cell.fill = accent_fill
    cash_cell.border = border
    
    for col in range(1, 7):
        ws3.column_dimensions[get_column_letter(col)].width = 22
    ws3.column_dimensions['A'].width = 30
    
    # ===== SHEET 4: WAITER SUMMARY =====
    ws4 = wb.create_sheet("Waiter Summary")
    waiter_users = User.objects.filter(profile__role='waiter', is_active=True).order_by('first_name')
    
    ws4.merge_cells('A1:C1')
    ws4['A1'].value = f"Waiter Summary - {label}"
    ws4['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws4['A1'].fill = primary_fill
    ws4['A1'].alignment = Alignment(horizontal='center')
    
    headers4 = ['Waiter', 'Number of Orders', 'Total Collection (TZS)']
    for col, header in enumerate(headers4, 1):
        cell = ws4.cell(row=3, column=col, value=header)
        cell.font = white_font
        cell.fill = secondary_fill
        cell.border = border
    
    for r, wu in enumerate(waiter_users, 4):
        waiter_orders = [o for o in orders if o.waiter_user == wu]
        display = wu.get_full_name() or wu.username
        ws4.cell(row=r, column=1, value=display).border = border
        ws4.cell(row=r, column=2, value=len(waiter_orders)).border = border
        ws4.cell(row=r, column=3, value=sum(float(o.total_price) for o in waiter_orders)).border = border
    
    for col in range(1, 4):
        ws4.column_dimensions[get_column_letter(col)].width = 25
    
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"DEEM_Report_{date_from}_to_{date_to}.xlsx" if date_from != date_to else f"DEEM_Report_{date_from}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
        # ===== SHEET 5: CANCELLED ORDERS =====
    cancelled_orders = Order.objects.filter(
        is_cancelled=True,
        created_at__date__gte=date_from,
        created_at__date__lte=date_to,
    ).select_related('menu_item', 'waiter_user', 'created_by', 'cancelled_by').order_by('-cancelled_at')
    
    ws5 = wb.create_sheet("Cancelled Orders")
    
    ws5.merge_cells('A1:H1')
    ws5['A1'].value = f"DEEM Restaurant - Cancelled Orders ({label})"
    ws5['A1'].font = Font(size=14, bold=True, color="FFFFFF")
    ws5['A1'].fill = danger_fill
    ws5['A1'].alignment = Alignment(horizontal='center')
    
    headers5 = ['Reference', 'Order Date', 'Menu Items', 'Qty', 'Amount (TZS)', 
                'Waiter', 'Cancelled At', 'Cancelled By', 'Reason']
    for col, header in enumerate(headers5, 1):
        cell = ws5.cell(row=3, column=col, value=header)
        cell.font = white_font
        cell.fill = danger_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center')
    
    cancelled_groups = defaultdict(list)
    for co in cancelled_orders:
        session_key = co.session_id if co.session_id else f"single_{co.id}"
        cancelled_groups[session_key].append(co)
    
    row5 = 4
    cancelled_total = 0
    for session_key, co_list in cancelled_groups.items():
        first_co = co_list[0]
        combined = OrderedDict()
        for o in co_list:
            key = o.menu_item.name
            combined[key] = combined.get(key, 0) + o.quantity
        
        items_str = ", ".join([f"{name} x{qty}" if qty > 1 else name for name, qty in combined.items()])
        total_qty = sum(o.quantity for o in co_list)
        total_amt = sum(float(o.total_price) for o in co_list)
        cancelled_total += total_amt
        
        ws5.cell(row=row5, column=1, value=f"#{first_co.group_id or first_co.id}").border = border
        ws5.cell(row=row5, column=2, value=first_co.created_at.strftime('%d/%m/%Y %H:%M')).border = border
        ws5.cell(row=row5, column=3, value=items_str).border = border
        ws5.cell(row=row5, column=4, value=total_qty).border = border
        ws5.cell(row=row5, column=5, value=total_amt).border = border
        ws5.cell(row=row5, column=6, value=first_co.get_waiter_display_name()).border = border
        ws5.cell(row=row5, column=7, value=first_co.cancelled_at.strftime('%d/%m/%Y %H:%M') if first_co.cancelled_at else '-').border = border
        ws5.cell(row=row5, column=8, value=first_co.cancelled_by.username if first_co.cancelled_by else '-').border = border
        ws5.cell(row=row5, column=9, value=first_co.cancellation_reason or '-').border = border
        
        row5 += 1
    
    if cancelled_groups:
        ws5.merge_cells(f'A{row5}:D{row5}')
        ws5.cell(row=row5, column=1, value="TOTAL CANCELLED (Excluded from Sales)").font = Font(bold=True, color="FFFFFF")
        ws5.cell(row=row5, column=1).fill = danger_fill
        ws5.cell(row=row5, column=1).border = border
        ws5.cell(row=row5, column=5, value=cancelled_total).font = Font(bold=True, color="FFFFFF")
        ws5.cell(row=row5, column=5).fill = danger_fill
        ws5.cell(row=row5, column=5).border = border
    
    for col in range(1, 10):
        ws5.column_dimensions[get_column_letter(col)].width = 20
    ws5.column_dimensions['C'].width = 40
    ws5.column_dimensions['I'].width = 40
    
    wb.save(response)
    return response


# ============================================================
# USER MANAGEMENT
# ============================================================

@admin_required
def user_management(request):
    users = User.objects.all().select_related('profile').order_by('-date_joined')
    users_data = []
    for u in users:
        try:
            profile = u.profile
        except UserProfile.DoesNotExist:
            profile = UserProfile.objects.create(user=u)
        users_data.append({'user': u, 'profile': profile})
    
    context = {
        'users_data': users_data,
        'total_users': users.count(),
        'admin_count': UserProfile.objects.filter(role='admin').count(),
        'manager_count': UserProfile.objects.filter(role='manager').count(),
        'cashier_count': UserProfile.objects.filter(role='cashier').count(),
        'waiter_count': UserProfile.objects.filter(role='waiter').count(),
        'active_count': User.objects.filter(is_active=True).count(),
    }
    return render(request, 'sales/user_management.html', context)


@admin_required
def add_user(request):
    """Add a new user - Admin only"""
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '')
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        role = request.POST.get('role', 'cashier')
        phone = request.POST.get('phone', '').strip()
        
        if not username or not password:
            messages.error(request, 'Username and password are required.')
            return redirect('user_management')
        
        if User.objects.filter(username=username).exists():
            messages.error(request, f'Username "{username}" already exists.')
            return redirect('user_management')
        
        if len(password) < 6:
            messages.error(request, 'Password must be at least 6 characters.')
            return redirect('user_management')
        
        try:
            with transaction.atomic():
                user = User.objects.create_user(
                    username=username,
                    email=email,
                    password=password,
                    first_name=first_name,
                    last_name=last_name,
                )
                
                if role == 'admin':
                    user.is_staff = True
                
                user.save()
                
                profile = user.profile
                profile.role = role
                profile.phone = phone
                profile.save()
                
                messages.success(request, f'User "{username}" created as {role.title()}!')
        except Exception as e:
            messages.error(request, f'Error creating user: {str(e)}')
    
    return redirect('user_management')


@admin_required
def edit_user(request, user_id):
    """Edit existing user - Admin only"""
    if request.method == 'POST':
        user = get_object_or_404(User, id=user_id)
        
        if user == request.user and request.POST.get('role') != 'admin':
            messages.error(request, 'You cannot change your own role.')
            return redirect('user_management')
        
        try:
            with transaction.atomic():
                from .audit import serialize_instance, log_action
                data_before = serialize_instance(user)
                
                user.email = request.POST.get('email', '').strip()
                user.first_name = request.POST.get('first_name', '').strip()
                user.last_name = request.POST.get('last_name', '').strip()
                
                role = request.POST.get('role', 'cashier')
                is_active = request.POST.get('is_active') == 'on'
                
                if role == 'admin':
                    user.is_staff = True
                else:
                    user.is_staff = False
                
                user.is_active = is_active
                user.save()
                
                profile = user.profile
                profile.role = role
                profile.phone = request.POST.get('phone', '').strip()
                profile.is_active = is_active
                profile.save()
                
                new_password = request.POST.get('new_password', '')
                if new_password and len(new_password) >= 6:
                    user.set_password(new_password)
                    user.save()
                
                log_action(request, 'update', user, data_before=data_before,
                           notes=f'Role: {role}, Active: {is_active}')
                
                messages.success(request, f'User "{user.username}" updated successfully!')
        except Exception as e:
            messages.error(request, f'Error updating user: {str(e)}')
    
    return redirect('user_management')


@admin_required
def toggle_user_status(request, user_id):
    if request.method == 'POST':
        user = get_object_or_404(User, id=user_id)
        
        if user == request.user:
            messages.error(request, 'You cannot deactivate your own account.')
            return redirect('user_management')
        
        user.is_active = not user.is_active
        user.save()
        
        try:
            user.profile.is_active = user.is_active
            user.profile.save()
        except UserProfile.DoesNotExist:
            pass
        
        status = "activated" if user.is_active else "deactivated"
        messages.success(request, f'User "{user.username}" {status}.')
    
    return redirect('user_management')


@admin_required
def delete_user(request, user_id):
    if request.method == 'POST':
        user = get_object_or_404(User, id=user_id)
        
        if user == request.user:
            messages.error(request, 'You cannot delete your own account.')
            return redirect('user_management')
        
        from .audit import serialize_instance, log_action
        data_before = serialize_instance(user)
        log_action(request, 'delete', user, data_before=data_before)
        
        username = user.username
        user.delete()
        messages.success(request, f'User "{username}" deleted successfully.')
    
    return redirect('user_management')


@admin_required
def reset_password(request, user_id):
    if request.method == 'POST':
        user = get_object_or_404(User, id=user_id)
        new_password = request.POST.get('new_password', '')
        
        if len(new_password) < 6:
            messages.error(request, 'Password must be at least 6 characters.')
            return redirect('user_management')
        
        user.set_password(new_password)
        user.save()
        messages.success(request, f'Password reset for "{user.username}".')
    
    return redirect('user_management')


# ============================================================
# MENU MANAGEMENT
# ============================================================

@reports_required
def menu_management(request):
    categories = MenuCategory.objects.all().prefetch_related('items')
    menu_items = MenuItem.objects.all().select_related('category').order_by('category__display_order', 'display_order', 'name')
    
    context = {
        'categories': categories,
        'menu_items': menu_items,
        'total_categories': categories.count(),
        'total_items': menu_items.count(),
        'available_items': menu_items.filter(is_available=True).count(),
    }
    return render(request, 'sales/menu_management.html', context)


@reports_required
def add_category(request):
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        display_order = request.POST.get('display_order', 0)
        
        if not name:
            messages.error(request, 'Category name is required.')
            return redirect('menu_management')
        
        if MenuCategory.objects.filter(name__iexact=name).exists():
            messages.error(request, f'Category "{name}" already exists.')
            return redirect('menu_management')
        
        MenuCategory.objects.create(name=name, display_order=int(display_order) if display_order else 0)
        messages.success(request, f'Category "{name}" created successfully!')
    
    return redirect('menu_management')


@reports_required
def edit_category(request, category_id):
    if request.method == 'POST':
        category = get_object_or_404(MenuCategory, id=category_id)
        name = request.POST.get('name', '').strip()
        display_order = request.POST.get('display_order', 0)
        
        if not name:
            messages.error(request, 'Category name is required.')
            return redirect('menu_management')
        
        category.name = name
        category.display_order = int(display_order) if display_order else 0
        category.save()
        messages.success(request, 'Category updated successfully!')
    
    return redirect('menu_management')


@reports_required
def delete_category(request, category_id):
    if request.method == 'POST':
        category = get_object_or_404(MenuCategory, id=category_id)
        
        if category.items.exists():
            messages.error(request, f'Cannot delete "{category.name}" — it has {category.items.count()} menu items.')
            return redirect('menu_management')
        
        name = category.name
        category.delete()
        messages.success(request, f'Category "{name}" deleted.')
    
    return redirect('menu_management')


@reports_required
def add_menu_item(request):
    if request.method == 'POST':
        name = request.POST.get('name', '').strip()
        category_id = request.POST.get('category')
        price = request.POST.get('price', '0')
        display_order = request.POST.get('display_order', 0)
        is_available = request.POST.get('is_available') == 'on'
        
        if not name or not category_id:
            messages.error(request, 'Menu item name and category are required.')
            return redirect('menu_management')
        
        try:
            category = MenuCategory.objects.get(id=category_id)
            MenuItem.objects.create(
                name=name,
                category=category,
                price=float(price),
                display_order=int(display_order) if display_order else 0,
                is_available=is_available
            )
            messages.success(request, f'Menu item "{name}" added successfully!')
        except MenuCategory.DoesNotExist:
            messages.error(request, 'Selected category not found.')
        except Exception as e:
            messages.error(request, f'Error adding menu item: {str(e)}')
    
    return redirect('menu_management')


@reports_required
def edit_menu_item(request, item_id):
    if request.method == 'POST':
        item = get_object_or_404(MenuItem, id=item_id)
        try:
            from .audit import serialize_instance, log_action
            data_before = serialize_instance(item)
            
            item.name = request.POST.get('name', '').strip()
            item.category = MenuCategory.objects.get(id=request.POST.get('category'))
            item.price = float(request.POST.get('price', '0'))
            item.display_order = int(request.POST.get('display_order', 0) or 0)
            item.is_available = request.POST.get('is_available') == 'on'
            item.save()
            
            log_action(request, 'update', item, data_before=data_before)
            messages.success(request, f'Menu item updated successfully!')
        except Exception as e:
            messages.error(request, f'Error updating menu item: {str(e)}')
    
    return redirect('menu_management')


@reports_required
def delete_menu_item(request, item_id):
    if request.method == 'POST':
        item = get_object_or_404(MenuItem, id=item_id)
        
        from .audit import serialize_instance, log_action
        data_before = serialize_instance(item)
        log_action(request, 'delete', item, data_before=data_before)
        
        name = item.name
        item.delete()
        messages.success(request, f'Menu item "{name}" deleted.')
    
    return redirect('menu_management')


@reports_required
def toggle_item_availability(request, item_id):
    if request.method == 'POST':
        item = get_object_or_404(MenuItem, id=item_id)
        item.is_available = not item.is_available
        item.save()
        status = "available" if item.is_available else "unavailable"
        messages.success(request, f'"{item.name}" marked as {status}.')
    
    return redirect('menu_management')


# ============================================================
# EXPENSES
# ============================================================

@expenses_required
def expenses_view(request):
    today = timezone.now().date()
    selected_date = request.GET.get('date', today.strftime('%Y-%m-%d'))
    
    try:
        selected_date = datetime.strptime(selected_date, '%Y-%m-%d').date()
    except:
        selected_date = today
    
    expenses = Expense.objects.filter(date=selected_date).select_related('recorded_by')
    today_expenses = Expense.objects.filter(date=today)
    today_expense_total = sum(float(e.amount) for e in today_expenses)
    today_orders = Order.objects.filter(created_at__date=today, is_cancelled=False)
    today_sales = sum(float(o.total_price) for o in today_orders)
    
    context = {
        'expenses': expenses,
        'selected_date': selected_date,
        'categories': Expense.CATEGORY_CHOICES,
        'today': today,
        'selected_total': sum(float(e.amount) for e in expenses),
        'today_expense_total': today_expense_total,
        'today_sales': today_sales,
        'today_balance': today_sales - today_expense_total,
    }
    return render(request, 'sales/expenses.html', context)


@expenses_required
def add_expense(request):
    if request.method == 'POST':
        try:
            category = request.POST.get('category', 'other')
            amount = request.POST.get('amount', '0')
            description = request.POST.get('description', '').strip()
            date_str = request.POST.get('date', '')
            receipt_ref = request.POST.get('receipt_ref', '').strip()
            
            if not description:
                messages.error(request, 'Description is required.')
                return redirect('expenses_view')
            
            expense_date = datetime.strptime(date_str, '%Y-%m-%d').date() if date_str else timezone.now().date()
            
            Expense.objects.create(
                category=category,
                amount=float(amount),
                description=description,
                date=expense_date,
                receipt_ref=receipt_ref if receipt_ref else None,
                recorded_by=request.user
            )
            messages.success(request, f'Expense of {float(amount):,.0f} TZS recorded successfully!')
        except Exception as e:
            messages.error(request, f'Error recording expense: {str(e)}')
    
    return redirect('expenses_view')


@expenses_required
def edit_expense(request, expense_id):
    if request.method == 'POST':
        expense = get_object_or_404(Expense, id=expense_id)
        try:
            from .audit import serialize_instance, log_action
            data_before = serialize_instance(expense)
            
            expense.category = request.POST.get('category', 'other')
            expense.amount = float(request.POST.get('amount', '0'))
            expense.description = request.POST.get('description', '').strip()
            date_str = request.POST.get('date', '')
            if date_str:
                expense.date = datetime.strptime(date_str, '%Y-%m-%d').date()
            expense.receipt_ref = request.POST.get('receipt_ref', '').strip() or None
            expense.save()
            
            log_action(request, 'update', expense, data_before=data_before)
            messages.success(request, 'Expense updated successfully!')
        except Exception as e:
            messages.error(request, f'Error updating expense: {str(e)}')
    
    return redirect('expenses_view')


@expenses_required
def delete_expense(request, expense_id):
    if request.method == 'POST':
        expense = get_object_or_404(Expense, id=expense_id)
        
        from .audit import serialize_instance, log_action
        data_before = serialize_instance(expense)
        log_action(request, 'delete', expense, data_before=data_before)
        
        expense.delete()
        messages.success(request, 'Expense deleted.')
    
    return redirect('expenses_view')


@expenses_required
def get_expense(request, expense_id):
    if request.method == 'GET':
        expense = get_object_or_404(Expense, id=expense_id)
        return JsonResponse({
            'id': expense.id,
            'category': expense.category,
            'amount': float(expense.amount),
            'description': expense.description,
            'date': expense.date.strftime('%Y-%m-%d'),
            'receipt_ref': expense.receipt_ref or '',
        })


# ============================================================
# AUDIT LOG
# ============================================================

@admin_required
def audit_log_view(request):
    action_filter = request.GET.get('action', '')
    model_filter = request.GET.get('model', '')
    user_filter = request.GET.get('user', '')
    date_filter = request.GET.get('date', '')
    
    logs = AuditLog.objects.all().select_related('performed_by')
    
    if action_filter:
        logs = logs.filter(action=action_filter)
    if model_filter:
        logs = logs.filter(model_name=model_filter)
    if user_filter:
        logs = logs.filter(performed_by__username=user_filter)
    if date_filter:
        try:
            filter_date = datetime.strptime(date_filter, '%Y-%m-%d').date()
            logs = logs.filter(performed_at__date=filter_date)
        except:
            pass
    
    logs = logs[:500]
    
    context = {
        'logs': logs,
        'total_logs': AuditLog.objects.count(),
        'update_count': AuditLog.objects.filter(action='update').count(),
        'delete_count': AuditLog.objects.filter(action='delete').count(),
        'today_logs': AuditLog.objects.filter(performed_at__date=timezone.now().date()).count(),
        'action_filter': action_filter,
        'model_filter': model_filter,
        'user_filter': user_filter,
        'date_filter': date_filter,
        'today': timezone.now().date(),
    }
    return render(request, 'sales/audit_log.html', context)


@admin_required
def audit_log_detail(request, log_id):
    log = get_object_or_404(AuditLog, id=log_id)
    
    return JsonResponse({
        'id': log.id,
        'action': log.get_action_display(),
        'model_name': log.model_name,
        'object_id': log.object_id,
        'object_repr': log.object_repr,
        'performed_at': log.performed_at.strftime('%d/%m/%Y %H:%M:%S'),
        'performed_by': log.performed_by.username if log.performed_by else 'system',
        'ip_address': log.ip_address or '',
        'notes': log.notes or '',
        'data_before': log.data_before,
        'data_after': log.data_after,
        'changes': log.get_changed_fields(),
    })
    
@login_required
def get_orders_page(request):
    """AJAX endpoint — returns just the table HTML for a paginated tab"""
    from django.template.loader import render_to_string
    from django.core.paginator import Paginator
    from django.db.models import Sum
    
    tab = request.GET.get('tab', 'latest')
    page_num = request.GET.get('page', 1)
    
    user_role = get_user_role(request.user)
    is_waiter = user_role == 'waiter'
    today = timezone.now().date()
    
    base_orders = Order.objects.all()
    if is_waiter:
        base_orders = base_orders.filter(waiter_user=request.user)
    
    if tab == 'latest':
        today_orders = base_orders.filter(
            created_at__date=today,
            is_cancelled=False
        ).select_related('menu_item', 'waiter_user', 'created_by').order_by('-created_at')
        
        grouped = _build_grouped_list(today_orders, today)
        paginator = Paginator(grouped, 20)
        page = paginator.get_page(page_num)
        
        html = render_to_string('sales/_latest_table.html', {
            'grouped_orders': page.object_list,
            'latest_page': page,
        }, request=request)
    else:  # unpaid
        unpaid_orders = base_orders.filter(
            status='open',
            is_cancelled=False
        ).select_related('menu_item', 'waiter_user', 'created_by').order_by('-created_at')
        
        grouped = _build_grouped_list(unpaid_orders, today, is_unpaid=True)
        paginator = Paginator(grouped, 20)
        page = paginator.get_page(page_num)
        
        html = render_to_string('sales/_unpaid_table.html', {
            'unpaid_grouped': page.object_list,
            'unpaid_page': page,
        }, request=request)
    
    return JsonResponse({'html': html, 'page': page.number})


def _build_grouped_list(orders_qs, today, is_unpaid=False):
    """Shared logic to build the grouped order list"""
    session_groups = defaultdict(list)
    for order in orders_qs:
        session_key = order.session_id if order.session_id else f"single_{order.id}"
        session_groups[session_key].append(order)
    
    result = []
    for idx, (session_key, orders) in enumerate(session_groups.items(), 1):
        combined_items = OrderedDict()
        for order in orders:
            key = order.menu_item.id
            if key in combined_items:
                combined_items[key]['quantity'] += order.quantity
                combined_items[key]['total_price'] += float(order.total_price)
                combined_items[key]['order_ids'].append(order.id)
            else:
                combined_items[key] = {
                    'menu_item': order.menu_item,
                    'quantity': order.quantity,
                    'unit_price': float(order.unit_price),
                    'total_price': float(order.total_price),
                    'category': order.menu_item.category,
                    'order_ids': [order.id],
                }
        
        combined_list = list(combined_items.values())
        first = orders[0]
        
        all_closed = all(o.status == 'closed' for o in orders)
        any_closed = any(o.status == 'closed' for o in orders)
        payment_status = 'closed' if all_closed else ('partial' if any_closed else 'open')
        
        group = {
            'order_number': idx,
            'session_id': session_key,
            'combined_items': combined_list,
            'created_at': first.created_at,
            'waiter_user': first.waiter_user,
            'waiter_name': first.get_waiter_display_name(),
            'created_by': first.created_by,
            'location_type': first.location_type,
            'location_detail': first.location_detail,
            'location_display': first.get_full_location_display(),
            'notes': first.notes or '',
            'total_quantity': sum(i['quantity'] for i in combined_list),
            'total_amount': sum(i['total_price'] for i in combined_list),
            'item_count': len(combined_list),
            'is_group': len(orders) > 1 or len(combined_list) > 1,
            'first_order_id': orders[-1].group_id or orders[-1].id,
            'last_order_id': orders[0].id,
            'payment_status': payment_status,
            'is_printed': first.is_printed,
            'is_cancelled': first.is_cancelled,
            'is_locked': first.is_locked,
            'cancellation_reason': first.cancellation_reason if first.is_cancelled else '',
        }
        
        if is_unpaid:
            group['days_old'] = (today - first.created_at.date()).days
        
        result.append(group)
    
    result.sort(key=lambda x: x['created_at'], reverse=True)
    for idx, g in enumerate(result, 1):
        g['order_number'] = idx
    
    return result

# ============================================================
# PRINT SYSTEM
# ============================================================

@login_required
def print_ticket(request, session_id):
    """Create a print job for kitchen or customer ticket"""
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    
    try:
        print_type = request.POST.get('print_type', 'kitchen')
        if print_type not in ['kitchen', 'receipt']:
            return JsonResponse({'error': 'Invalid print type'}, status=400)
        
        orders = Order.objects.filter(session_id=session_id).select_related(
            'menu_item', 'menu_item__category', 'waiter_user'
        )
        
        if not orders.exists():
            return JsonResponse({'error': 'Order not found'}, status=404)
        
        if orders.filter(is_cancelled=True).exists():
            return JsonResponse({'error': 'Cannot print a cancelled order'}, status=400)
        
        first = orders.first()
        
        # Build the ticket text
        if print_type == 'kitchen':
            ticket_text = build_kitchen_ticket(orders, first)
        else:
            ticket_text = build_customer_receipt(orders, first)
        
        job = PrintJob.objects.create(
            session_id=session_id,
            reference=first.reference or f"ORD-{first.id}",
            print_type=print_type,
            status='pending',
            requested_by=request.user,
            ticket_text=ticket_text,
        )
        
        # Lock the order after first print (any type)
        if not first.is_printed:
            orders.update(is_printed=True, printed_at=timezone.now())
        
        return JsonResponse({
            'success': True,
            'job_id': job.id,
            'reference': job.reference,
            'print_type': print_type,
            'status': job.status,
        })
    
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)


def build_kitchen_ticket(orders, first_order):
    """Kitchen ticket — no prices, focus on what to cook"""
    from io import StringIO
    from collections import OrderedDict
    
    out = StringIO()
    W = 42  # 80mm width
    
    combined = OrderedDict()
    for o in orders:
        key = o.menu_item.id
        if key in combined:
            combined[key]['quantity'] += o.quantity
        else:
            combined[key] = {
                'name': o.menu_item.name,
                'quantity': o.quantity,
            }
    
    total_qty = sum(i['quantity'] for i in combined.values())
    ref = first_order.reference or f"ORD-{first_order.id}"
    time_str = first_order.created_at.strftime('%H:%M')
    waiter = first_order.get_waiter_display_name()
    
    if first_order.location_type == 'other':
        location = first_order.location_detail or 'Other'
    elif first_order.location_detail:
        location = f"{first_order.get_location_type_display()} #{first_order.location_detail}"
    else:
        location = first_order.get_location_type_display()
    
    out.write('=' * W + '\n')
    out.write('KITCHEN TICKET'.center(W) + '\n')
    out.write(ref.center(W) + '\n')
    out.write('=' * W + '\n')
    out.write(f"Time: {time_str}  |  Waiter: {waiter}\n")
    out.write(f"Place: {location}  |  Items: {total_qty}\n")
    out.write('-' * W + '\n')
    out.write(f"{'QTY':>4}  ITEM\n")
    out.write('-' * W + '\n')
    
    for item in combined.values():
        out.write(f"{item['quantity']:>4}  {item['name']}\n")
    
    out.write('-' * W + '\n')
    
    if first_order.notes:
        out.write('NOTES:\n')
        for line in first_order.notes.split('\n'):
            out.write(f"  {line}\n")
        out.write('-' * W + '\n')
    
    out.write('\n\n\n')
    return out.getvalue()


def build_customer_receipt(orders, first_order):
    """Customer receipt — with prices and total"""
    from io import StringIO
    from collections import OrderedDict
    
    out = StringIO()
    W = 42
    
    combined = OrderedDict()
    for o in orders:
        key = o.menu_item.id
        if key in combined:
            combined[key]['quantity'] += o.quantity
            combined[key]['total'] += float(o.total_price)
        else:
            combined[key] = {
                'name': o.menu_item.name,
                'quantity': o.quantity,
                'unit_price': float(o.unit_price),
                'total': float(o.total_price),
            }
    
    total_qty = sum(i['quantity'] for i in combined.values())
    total_amount = sum(i['total'] for i in combined.values())
    ref = first_order.reference or f"ORD-{first_order.id}"
    time_str = first_order.created_at.strftime('%H:%M')
    waiter = first_order.get_waiter_display_name()
    
    if first_order.location_type == 'other':
        location = first_order.location_detail or 'Other'
    elif first_order.location_detail:
        location = f"{first_order.get_location_type_display()} #{first_order.location_detail}"
    else:
        location = first_order.get_location_type_display()
    
    out.write('=' * W + '\n')
    out.write('DEEM RESTAURANT'.center(W) + '\n')
    out.write('Dar es Salaam'.center(W) + '\n')
    out.write('+255 665 304 561'.center(W) + '\n')
    out.write('=' * W + '\n')
    out.write(f"Receipt: {ref}\n")
    out.write(f"Time: {time_str}  |  Waiter: {waiter}\n")
    out.write(f"Place: {location}\n")
    out.write('-' * W + '\n')
    out.write(f"{'ITEM':<20} {'QTY':>4} {'AMOUNT':>12}\n")
    out.write('-' * W + '\n')
    
    for item in combined.values():
        name = item['name'][:20]
        out.write(f"{name:<20} {item['quantity']:>4} {item['total']:>12,.0f}\n")
    
    out.write('-' * W + '\n')
    out.write(f"{'TOTAL':<20} {total_qty:>4} {total_amount:>12,.0f}\n")
    out.write('=' * W + '\n')
    out.write('Asante! Karibu tena'.center(W) + '\n')
    out.write('=' * W + '\n')
    out.write(f"Printed: {timezone.now().strftime('%H:%M:%S')}".rjust(W) + '\n')
    out.write('\n\n\n')
    return out.getvalue()


@login_required
def print_job_status(request, job_id):
    """Poll status of a print job"""
    try:
        job = PrintJob.objects.get(id=job_id)
        return JsonResponse({
            'id': job.id,
            'status': job.status,
            'reference': job.reference,
            'print_type': job.print_type,
            'error': job.error_message,
            'printed_at': job.printed_at.strftime('%H:%M:%S') if job.printed_at else None,
        })
    except PrintJob.DoesNotExist:
        return JsonResponse({'error': 'Job not found'}, status=404)


def _check_agent_key(request):
    """Verify the local print agent's secret"""
    from django.conf import settings
    key = request.headers.get('X-Print-Agent-Key', '')
    return key == getattr(settings, 'PRINT_AGENT_KEY', '')


def print_agent_pending(request):
    """Agent polls for pending jobs"""
    if not _check_agent_key(request):
        return JsonResponse({'error': 'Unauthorized'}, status=401)
    
    jobs = PrintJob.objects.filter(status='pending').order_by('requested_at')[:10]
    
    payload = []
    for job in jobs:
        payload.append({
            'job_id': job.id,
            'reference': job.reference,
            'print_type': job.print_type,
            'ticket_text': job.ticket_text or '',
        })
        job.status = 'sent'
        job.save(update_fields=['status'])
    
    return JsonResponse({'jobs': payload})


def print_agent_report(request, job_id):
    """Agent reports print result"""
    if not _check_agent_key(request):
        return JsonResponse({'error': 'Unauthorized'}, status=401)
    
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    
    try:
        job = PrintJob.objects.get(id=job_id)
        data = json.loads(request.body)
        status = data.get('status', 'failed')
        error = data.get('error', '')
        
        if status == 'printed':
            job.status = 'printed'
            job.printed_at = timezone.now()
            job.error_message = None
        else:
            job.status = 'failed'
            job.error_message = error[:255]
            job.retry_count += 1
        
        job.save()
        return JsonResponse({'success': True})
    
    except PrintJob.DoesNotExist:
        return JsonResponse({'error': 'Job not found'}, status=404)
    
@login_required
def cancel_order_group(request, session_id):
    """Cancel an entire order group — irreversible, requires reason"""
    if request.method != 'POST':
        return JsonResponse({'error': 'POST required'}, status=405)
    
    try:
        reason = request.POST.get('reason', '').strip()
        
        if not reason or len(reason) < 5:
            return JsonResponse({'error': 'Cancellation reason required (min 5 characters)'}, status=400)
        
        orders = Order.objects.filter(session_id=session_id)
        if not orders.exists():
            return JsonResponse({'error': 'Order not found'}, status=404)
        
        if orders.filter(is_cancelled=True).exists():
            return JsonResponse({'error': 'Order already cancelled'}, status=400)
        
        # NEW: Cannot cancel PAID order
        all_paid = all(o.status == 'closed' for o in orders)
        if all_paid:
            return JsonResponse({
                'error': 'Cannot cancel a PAID order. Mark it as UNPAID first.'
            }, status=400)
        
        orders.update(
            status='cancelled',
            is_cancelled=True,
            cancelled_at=timezone.now(),
            cancelled_by=request.user,
            cancellation_reason=reason,
        )
        
        # Log to audit
        from .audit import log_action
        first = orders.first()
        log_action(request, 'update', first, data_before={}, 
                   notes=f'CANCELLED: {reason}')
        
        return JsonResponse({
            'success': True,
            'message': f'Order cancelled. Reason: {reason}',
        })
    
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)
    

# Temporary for printing testing purpose
@login_required
def preview_ticket(request, session_id):
    """Preview the ticket text in browser (for testing without printer)"""
    print_type = request.GET.get('type', 'kitchen')
    orders = Order.objects.filter(session_id=session_id).select_related('menu_item', 'menu_item__category', 'waiter_user')
    
    if not orders.exists():
        return HttpResponse('Order not found', status=404)
    
    first = orders.first()
    
    if print_type == 'kitchen':
        text = build_kitchen_ticket(orders, first)
    else:
        text = build_customer_receipt(orders, first)
    
    return HttpResponse(text, content_type='text/plain; charset=utf-8')